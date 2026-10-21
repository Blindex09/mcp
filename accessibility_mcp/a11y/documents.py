"""Geracao de DOCX semantico e PDF marcado, sem chamar modelo nem executar HTML do cliente."""

from __future__ import annotations

import asyncio
import base64
import io
import json
import re
import uuid
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from playwright.async_api import async_playwright
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    BooleanObject,
    ContentStream,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    TextStringObject,
)

from . import files, learn
from .document_schema import Block, DocumentSpec
from .document_validation import validate
from .provisioning import provisioner_for

Formats = Literal["docx", "pdf", "both"]
MIMES = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "pdf": "application/pdf"}


def folder() -> Path:
    return learn.store_path().resolve() / "documents"


def artifact_path(document_id: str, filename: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", document_id) or filename not in ("document.docx", "document.pdf"):
        raise ValueError("Identificador ou nome de documento invalido.")
    root = folder()
    path = (root / document_id / filename).resolve()
    if not path.is_relative_to(root):
        raise ValueError("Caminho fora da pasta de documentos.")
    return path


def resolve_input(path: str) -> Path:
    candidate = Path(path).expanduser().resolve()
    if candidate.is_relative_to(folder()):
        relative = candidate.relative_to(folder())
        if len(relative.parts) != 2:
            raise ValueError("Caminho de documento invalido.")
        return artifact_path(*relative.parts)
    return files.resolve(path)


def _images(spec: DocumentSpec) -> dict[str, bytes]:
    result: dict[str, bytes] = {}
    for block in spec.blocks:
        if block.kind == "image":
            path = files.resolve(block.source)
            if path.suffix.lower() not in (".png", ".jpg", ".jpeg"):
                raise ValueError("Imagem deve ser PNG ou JPEG local, dentro da pasta escolhida.")
            with Image.open(path) as image:
                image.load()
                stream = io.BytesIO()
                image.convert("RGB").save(stream, format="PNG")
                result[block.source] = stream.getvalue()
    return result


def _all_blocks(spec: DocumentSpec) -> list[Block]:
    blocks = list(spec.blocks)
    if spec.assessments:
        blocks.append(Block(kind="heading", text="Avaliação por critério", level=2))
        for assessment in spec.assessments:
            blocks.append(Block(kind="heading", text=f"WCAG {assessment.criterion}", level=3))
            blocks.append(Block(kind="paragraph", text=f"Resultado: {assessment.status}. Cobertura declarada: {assessment.coverage}. {assessment.remarks}"))
            if assessment.evidence:
                blocks.append(Block(kind="list", items=assessment.evidence))
        blocks.append(Block(kind="paragraph", text="Critérios não listados não foram avaliados neste documento. A avaliação se limita ao escopo e às evidências declarados."))
    return blocks


def _property(parent: Any, tag: str, **attrs: str) -> Any:
    element = OxmlElement(tag)
    for name, value in attrs.items():
        element.set(qn(name), value)
    parent.append(element)
    return element


def build_docx(spec: DocumentSpec, images: dict[str, bytes], path: Path) -> None:
    doc = Document()
    props = doc.core_properties
    props.title, props.subject, props.author = spec.title, spec.subject, spec.author
    props.language = spec.language
    props.comments = ""
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(8)
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr_default = defaults.find(qn("w:rPrDefault"))
    rpr = rpr_default.find(qn("w:rPr"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = _property(rpr, "w:lang")
    lang.set(qn("w:val"), spec.language)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        paragraph_properties = doc.styles[name].element.find(qn("w:pPr"))
        if paragraph_properties is not None:
            for border in list(paragraph_properties.findall(qn("w:pBdr"))):
                paragraph_properties.remove(border)
    doc.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Caption"].font.size = Pt(11)
    title = doc.add_paragraph(spec.title, "Title")
    _property(title._p.get_or_add_pPr(), "w:outlineLvl", **{"w:val": "0"})
    for block in _all_blocks(spec):
        if block.kind == "heading":
            doc.add_heading(block.text, level=block.level)
        elif block.kind == "paragraph":
            doc.add_paragraph(block.text)
        elif block.kind == "list":
            # Cada lista numerada tem sua propria instancia para recomecar em 1.
            style = "List Number" if block.ordered else "List Bullet"
            numbering = doc.part.numbering_part.element
            style_id = doc.styles[style].style_id
            abstract = numbering.xpath(f'./w:abstractNum[w:lvl/w:pStyle[@w:val="{style_id}"]]')
            num_id = None
            if abstract:
                num = numbering.add_num(int(abstract[0].get(qn("w:abstractNumId"))))
                num.add_lvlOverride(ilvl=0).add_startOverride(1)
                num_id = num.numId
            for item in block.items:
                paragraph = doc.add_paragraph(item, style)
                if num_id is not None:
                    numbering_properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    numbering_properties.get_or_add_ilvl().val = 0
                    numbering_properties.get_or_add_numId().val = num_id
        elif block.kind == "table":
            doc.add_paragraph(block.text, "Caption")
            table = doc.add_table(rows=1, cols=len(block.headers))
            table.style = "Table Grid"
            table.autofit = False
            _property(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader", **{"w:val": "true"})
            _property(table._tbl.tblPr, "w:tblCaption", **{"w:val": block.text})
            for cell, text in zip(table.rows[0].cells, block.headers):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
            for table_row in table.rows:
                _property(table_row._tr.get_or_add_trPr(), "w:cantSplit")
        elif block.kind == "image":
            with Image.open(io.BytesIO(images[block.source])) as image:
                width, height = image.size
            scale = min(12 / width, 20 / height)
            shape = doc.add_picture(io.BytesIO(images[block.source]), width=Cm(width * scale))
            shape._inline.docPr.set("descr", block.alt)
        elif block.kind == "link":
            paragraph = doc.add_paragraph()
            link = OxmlElement("w:hyperlink")
            rel = doc.part.relate_to(block.url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link.set(qn("r:id"), rel)
            run = _property(link, "w:r")
            properties = _property(run, "w:rPr")
            _property(properties, "w:color", **{"w:val": "000080"})
            _property(properties, "w:u", **{"w:val": "single"})
            _property(run, "w:t").text = block.text
            paragraph._p.append(link)
    doc.save(str(path))


def build_html(spec: DocumentSpec, images: dict[str, bytes]) -> str:
    body = [f"<h1>{escape(spec.title)}</h1>"]
    for block in _all_blocks(spec):
        text = escape(block.text)
        if block.kind == "heading":
            body.append(f"<h{block.level}>{text}</h{block.level}>")
        elif block.kind == "paragraph":
            body.append(f"<p>{text}</p>")
        elif block.kind == "list":
            tag = "ol" if block.ordered else "ul"
            body.append(f"<{tag}>" + "".join(f"<li>{escape(s)}</li>" for s in block.items) + f"</{tag}>")
        elif block.kind == "table":
            head = "".join(f'<th scope="col">{escape(h)}</th>' for h in block.headers)
            rows = "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in row) + "</tr>" for row in block.rows)
            body.append(f"<table><caption>{text}</caption><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
        elif block.kind == "image":
            encoded = base64.b64encode(images[block.source]).decode("ascii")
            body.append(f'<img src="data:image/png;base64,{encoded}" alt="{escape(block.alt, quote=True)}">')
        elif block.kind == "link":
            body.append(f'<p><a href="{escape(block.url, quote=True)}">{text}</a></p>')
    return f'''<!doctype html><html lang="{escape(spec.language, quote=True)}"><head><meta charset="utf-8">
<title>{escape(spec.title)}</title><meta name="author" content="{escape(spec.author, quote=True)}">
<style>@page {{size:A4;margin:20mm}} body {{font:11pt Arial,sans-serif;line-height:1.4;color:#000;background:#fff}}
h1 {{font-size:24pt}} h2 {{font-size:18pt}} h3 {{font-size:14pt}} h1,h2,h3,h4,h5,h6 {{break-after:avoid}}
p,li,td,th {{overflow-wrap:anywhere}} p {{orphans:3;widows:3}} table {{width:100%;border-collapse:collapse;table-layout:fixed}}
caption {{text-align:left;font-weight:bold;margin:8pt 0}} th,td {{border:1px solid #555;padding:6pt;text-align:left;vertical-align:top}}
thead {{display:table-header-group}} tr {{break-inside:avoid}} img {{max-width:100%;max-height:220mm;object-fit:contain}}
a {{color:#000080;text-decoration:underline}}</style></head><body>{''.join(body)}</body></html>'''


async def build_pdf(spec: DocumentSpec, images: dict[str, bytes], path: Path) -> None:
    await provisioner_for("chromium").ensure()
    async with async_playwright() as pw:
        browser = await pw.chromium.launch()
        try:
            page = await browser.new_page(java_script_enabled=False)
            await page.route("**/*", lambda route: route.abort())
            await page.set_content(build_html(spec, images))
            links = await page.locator("a").evaluate_all("anchors => Object.fromEntries(anchors.map(a => [a.href, a.textContent]))")
            await page.pdf(path=str(path), tagged=True, outline=True, prefer_css_page_size=True, print_background=True)
        finally:
            await browser.close()
    await asyncio.to_thread(_finish_pdf, spec, path, links)


def _finish_pdf(spec: DocumentSpec, path: Path, links: dict[str, str]) -> None:
    """Completa lacunas objetivas do exportador Chromium, preservando a arvore de tags.

    O Chromium deixa o fundo da pagina sem Artifact, links sem Contents e nao escreve XMP.
    So marcamos pintura grafica sem MCID como decoracao; texto nunca e descartado como artefato.
    """
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(path))
    painting = {b"f", b"F", b"f*", b"S", b"s", b"B", b"B*", b"b", b"b*", b"sh"}
    for page in writer.pages:
        page[NameObject("/Tabs")] = NameObject("/S")
        stream = ContentStream(page["/Contents"], writer)
        operations, depth = [], 0
        for operands, operator in stream.operations:
            if operator in (b"BDC", b"BMC"):
                depth += 1
            if depth == 0 and operator in painting:
                operations.extend([([NameObject("/Artifact")], b"BMC"), (operands, operator), ([], b"EMC")])
            else:
                operations.append((operands, operator))
            if operator == b"EMC":
                depth -= 1
        stream.operations = operations
        page[NameObject("/Contents")] = writer._add_object(stream)
        for ref in page.get("/Annots", []):
            annot = ref.get_object()
            action = annot.get("/A")
            action = action.get_object() if action is not None else {}
            uri = str(action.get("/URI", ""))
            if annot.get("/Subtype") == "/Link" and uri in links:
                annot[NameObject("/Contents")] = TextStringObject(links[uri])
    writer.root_object[NameObject("/ViewerPreferences")] = DictionaryObject({NameObject("/DisplayDocTitle"): BooleanObject(True)})
    metadata = DecodedStreamObject()
    metadata.set_data(f'''<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>
<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">
<rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:pdfuaid="http://www.aiim.org/pdfua/ns/id/">
<pdfuaid:part>1</pdfuaid:part><dc:title><rdf:Alt><rdf:li xml:lang="x-default">{escape(spec.title)}</rdf:li></rdf:Alt></dc:title>
<dc:language><rdf:Bag><rdf:li>{escape(spec.language)}</rdf:li></rdf:Bag></dc:language>
<dc:creator><rdf:Seq><rdf:li>{escape(spec.author)}</rdf:li></rdf:Seq></dc:creator>
</rdf:Description></rdf:RDF></x:xmpmeta><?xpacket end="w"?>'''.encode())
    metadata[NameObject("/Type")] = NameObject("/Metadata")
    metadata[NameObject("/Subtype")] = NameObject("/XML")
    writer.root_object[NameObject("/Metadata")] = writer._add_object(metadata)
    writer.generate_file_identifiers()
    buffer = io.BytesIO()
    writer.write(buffer)
    path.write_bytes(buffer.getvalue())


async def create(spec: DocumentSpec, formats: Formats = "both") -> dict[str, Any]:
    if formats not in ("docx", "pdf", "both"):
        raise ValueError("Formato deve ser docx, pdf ou both.")
    images = _images(spec)
    document_id = uuid.uuid4().hex
    target = folder() / document_id
    target.mkdir(parents=True, exist_ok=False)
    outputs: list[dict[str, Any]] = []
    for kind in ("docx", "pdf") if formats == "both" else (formats,):
        path = artifact_path(document_id, f"document.{kind}")
        if kind == "docx":
            await asyncio.to_thread(build_docx, spec, images, path)
        else:
            await build_pdf(spec, images, path)
        validation = await validate(path)
        outputs.append({"path": str(path), "uri": f"a11y://document/{document_id}/document.{kind}",
                        "mime_type": MIMES[kind], "bytes": path.stat().st_size, "validation": validation})
    (target / "validation.json").write_text(json.dumps(outputs, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"document_id": document_id, "files": outputs, "created": datetime.now(timezone.utc).isoformat(),
            "note": "Arquivos gerados. A validacao automatica nao substitui revisao visual, semantica e com leitor de tela."}


def report_spec(record: dict[str, Any], title: str, language: str = "pt-BR") -> DocumentSpec:
    if language not in ("pt", "pt-BR", "pt-PT"):
        raise ValueError("O relatorio automatico e redigido em portugues. Para outro idioma, escreva o conteudo e use a11y_create_document com o idioma correspondente.")
    blocks = [Block(kind="heading", text="Escopo da avaliação"),
              Block(kind="paragraph", text=f"Alvo: {record.get('url', '')}. Objetivo: {record.get('goal') or 'Não informado'}. Persona: {record.get('persona', 'Não informada')}. Data: {record.get('created') or datetime.now(timezone.utc).isoformat()}."),
              Block(kind="heading", text="Achados e evidências")]
    findings = record.get("findings", [])
    if not findings:
        blocks.append(Block(kind="paragraph", text="Nenhum achado registrado. Isso não comprova conformidade nem cobertura dos critérios."))
    for i, finding in enumerate(findings, 1):
        blocks.extend([Block(kind="heading", level=3, text=f"Achado {i} WCAG {finding.get('criterion') or 'não definido'}"),
                       Block(kind="paragraph", text=str(finding.get("claim") or "Sem descrição")),
                       Block(kind="paragraph", text=f"Resultado registrado: {finding.get('outcome', 'cantTell')}. Gravidade informada: {finding.get('severity', 'não informada')}. Passos de evidência: {finding.get('evidence_steps', [])}."),
                       Block(kind="paragraph", text=f"Elemento: {json.dumps(finding.get('element') or finding.get('target') or 'Não informado', ensure_ascii=False)}. Observações: {'; '.join(str(n) for n in finding.get('notes', [])) or 'Nenhuma observação adicional'}."),
                       Block(kind="paragraph", text="Verificação registrada: " + json.dumps(finding.get("verification") or "Sem verificação adicional registrada", ensure_ascii=False))])
    blocks.extend([Block(kind="heading", text="Cobertura e verificações pendentes"),
                   Block(kind="paragraph", text="A ausência de achados não é aprovação de critério. Os critérios sem avaliação completa e evidências permanecem não avaliados. Este relatório não certifica conformidade.")])
    coverage = record.get("coverage") or {}
    if coverage:
        interactive = coverage.get("interactive") or {}
        blocks.append(Block(kind="paragraph", text=f"Navegador: {coverage.get('browser', 'Não informado')}. Elementos descobertos: {interactive.get('discovered', 0)}; listados: {interactive.get('listed', 0)}; sondados por comportamento: {interactive.get('probed_by_behavior', 0)}. Ações realizadas: {coverage.get('actions_taken', 0)}."))
        blocks.append(Block(kind="paragraph", text="Sondas executadas: " + (", ".join(str(p) for p in coverage.get("probes_run", [])) or "Nenhuma registrada") + "."))
        if coverage.get("gaps"):
            blocks.append(Block(kind="list", items=[str(gap) for gap in coverage["gaps"]]))
    else:
        blocks.append(Block(kind="paragraph", text="Cobertura não registrada nesta execução."))
    return DocumentSpec(title=title, language=language, subject="Avaliação de acessibilidade com evidências", blocks=blocks)
