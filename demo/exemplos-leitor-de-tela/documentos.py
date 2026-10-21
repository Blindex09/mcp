"""Galeria de documentos reais: contratos pesquisados, pares e medidas reproduziveis.

Os defeitos sao deliberados e ficam somente nas copias didaticas, nunca no gerador MCP.
Uso: python demo/exemplos-leitor-de-tela/documentos.py
"""
from __future__ import annotations

import asyncio
import io
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(ROOT))

from accessibility_mcp.a11y import documents
from accessibility_mcp.a11y.document_schema import Block, DocumentSpec
from accessibility_mcp.a11y.document_validation import validate

CONTENT = ROOT / "accessibility_mcp/a11y/content"
DATA = json.loads((CONTENT / "assets/document-examples.json").read_text(encoding="utf-8"))
GUIDE = CONTENT / "references/document-behavior-matrix.md"
ISSUES = {
    "missing_heading_one": "título principal sem estrutura de título",
    "image_missing_alt": "figura informativa sem alternativa de texto",
    "table_missing_header": "tabela sem cabeçalho identificado",
    "missing_title": "metadado de título ausente",
    "missing_language": "idioma padrão ausente",
    "missing_tags": "estrutura de tags ausente",
    "page_without_extractable_text": "página sem texto extraível",
}


def guide_text() -> str:
    lines = ["# Document behavior matrix: Word DOCX and PDF examples", "",
             "Origem: este MCP. Contratos dos exemplos de documentos e uso das ferramentas deste servidor.",
             f"Fontes oficiais consultadas em {DATA['reviewed']}. Textos sintetizados; arquivos de terceiros não foram incorporados.", "",
             "## Como usar e o que pode afirmar", "",
             "Escolha os assuntos pelo significado da tarefa. A galeria está em demo/exemplos-leitor-de-tela/documentos.html.",
             "Para produzir arquivos, escreva o conteúdo e use a11y_create_document; para relatório dos achados, a11y_create_report. Valide o arquivo efetivo com a11y_validate_document.",
             "O MCP mede estrutura; o modelo do cliente julga significado, qualidade e cobertura. Não escolha casos por palavras-chave nem converta ausência de erros em conformidade.",
             "As técnicas W3C são formas de atender critérios, não uma lista de implementações obrigatórias. A galeria ensina casos delimitados; Bom significa que a propriedade exemplificada foi preservada, não certificação do arquivo.",
             "A inspeção DOCX verifica metadados, títulos, alternativas, cabeçalhos de tabela e links sem texto. PDF acrescenta tags, fontes, Unicode e texto extraível. Listas simuladas, propósito de links, sentido da alternativa e ordem lógica exigem julgamento e revisão no aplicativo.",
             "Uma imagem decorativa corretamente marcada não exige a mesma descrição de uma figura informativa. O validador DOCX atual não distingue decoração e pode pedir revisão; o gerador exige alt para as imagens que aceita.",
             "PDF/UA-1 e PDF/UA-2 são perfis distintos. O gerador atual é verificado automaticamente com veraPDF PDF/UA-1; não declare PDF/UA-2. Aprovação automática não comprova todos os requisitos nem WCAG.",
             "A sonda a11y_nvda_speech é para páginas na VM configurada, não abre Word ou leitor de PDF. Não invente fala literal nem declare teste real de documentos sem capturar o leitor no aplicativo. Registre versões, comandos, resultado e limitações. OCR do leitor não torna o arquivo acessível.", ""]
    for item in DATA["scenarios"]:
        lines.extend([f"## {item['title']}", "", f"Fazer: {item['do']}", f"Evitar: {item['dont']}",
                      f"Comportamento esperado, não fala capturada: {item['expected']}",
                      f"Verificação no aplicativo: {item['manual']}", f"Defeito demonstrado: {item['defect']}",
                      f"Medição automática: {item['detected'] or 'o validador estrutural pode não detectar este defeito; revisar o conteúdo e a estrutura específica' }.",
                      "Fontes: " + "; ".join(f"[{DATA['sources'][key]['title']}]({DATA['sources'][key]['url']})" for key in item["sources"]), ""])
    lines.extend(["## Limites complementares", "",
                  "Confira contraste e uso de cor sem depender apenas da cor; texto pequeno ou layout que não reflui podem impedir o uso com ampliação. Formulários precisam de nomes, instruções, estados e ordem de foco; o gerador atual não cria formulários interativos. Valide essas propriedades separadamente, sem afirmar que foram cobertas pelos pares desta galeria.",
                  "Não envie documentos privados a serviços externos sem autorização. Trabalhe no arquivo dentro da pasta escolhida pela pessoa e prefira corrigir a fonte, reexportar e conferir o resultado. Um DOCX correto não garante que qualquer exportação PDF preserve a estrutura.", ""])
    return "\n".join(lines)


def spec_for(slug: str) -> DocumentSpec:
    item = next(s for s in DATA["scenarios"] if s["id"] == slug)
    blocks = [Block(kind="paragraph", text="Exemplo didático para comparar a estrutura dos arquivos. Consulte a página da galeria antes de avaliar."),
              Block(kind="heading", text="Preparação do documento")]
    if slug == "imagens":
        blocks.append(Block(kind="image", source="grafico.png", alt="Participação nos testes: 12 pessoas em setembro e 18 em outubro."))
    elif slug == "tabelas":
        blocks.append(Block(kind="table", text="Participação nos testes", headers=["Mês", "Participantes"], rows=[["Setembro", "12"], ["Outubro", "18"]]))
    elif slug in ("listas", "ordem"):
        blocks.append(Block(kind="list", ordered=True, items=["Primeiro prepare o documento", "Depois revise o documento"]))
    elif slug == "links":
        blocks.append(Block(kind="link", text="Guia de acessibilidade do Word da Microsoft", url=DATA["sources"]["word"]["url"]))
    else:
        blocks.append(Block(kind="paragraph", text="Primeiro prepare o documento. Depois revise o documento. Confira a estrutura e leia o conteúdo no aplicativo de destino."))
    return DocumentSpec(title=item["title"], language="pt-BR", author="Projeto MCP de acessibilidade", blocks=blocks)


def image_bytes() -> bytes:
    picture = Image.new("RGB", (700, 220), "white")
    draw = ImageDraw.Draw(picture)
    font = ImageFont.load_default(size=24)
    for y, name, width, value in [(45, "Setembro", 240, "12"), (125, "Outubro", 360, "18")]:
        draw.text((12, y), name, font=font, fill="black")
        draw.rectangle((160, y, 160 + width, y + 38), fill="#173c63")
        draw.text((175 + width, y), value, font=font, fill="black")
    data = io.BytesIO()
    picture.save(data, "PNG")
    return data.getvalue()


def bad_docx(source: Path, target: Path, slug: str) -> None:
    doc = Document(source)
    if slug == "titulos":
        for paragraph in doc.paragraphs:
            if paragraph.style.name in ("Title", "Heading 2"):
                paragraph.style = doc.styles["Normal"]
                for outline in list(paragraph._p.xpath("./w:pPr/w:outlineLvl")):
                    outline.getparent().remove(outline)
                for run in paragraph.runs:
                    run.bold = True
    elif slug == "imagens":
        for shape in doc.inline_shapes:
            shape._inline.docPr.attrib.pop("descr", None)
    elif slug == "tabelas":
        for header in doc.element.xpath(".//w:tblHeader"):
            header.getparent().remove(header)
    elif slug == "listas":
        index = 0
        for paragraph in doc.paragraphs:
            if paragraph._p.xpath("./w:pPr/w:numPr"):
                index += 1
                text = paragraph.text
                paragraph.style = doc.styles["Normal"]
                for num in list(paragraph._p.xpath("./w:pPr/w:numPr")):
                    num.getparent().remove(num)
                paragraph.text = f"{index}. {text}"
    elif slug == "links":
        for text in doc.element.xpath(".//w:hyperlink//w:t"):
            text.text = "Clique aqui"
    elif slug == "metadados":
        doc.core_properties.title = doc.core_properties.language = ""
        for lang in doc.styles.element.xpath(".//w:lang"):
            lang.attrib.pop(qn("w:val"), None)
    doc.save(target)


def nodes(root: Any) -> list[Any]:
    found, seen = [], set()

    def walk(value: Any) -> None:
        value = value.get_object() if hasattr(value, "get_object") else value
        if isinstance(value, list):
            for child in value:
                walk(child)
        elif isinstance(value, dict) and id(value) not in seen:
            seen.add(id(value))
            if "/S" in value:
                found.append(value)
            if "/K" in value:
                walk(value["/K"])
    walk(root)
    return found


def bad_pdf(source: Path, target: Path, slug: str) -> None:
    writer = PdfWriter(clone_from=source)
    tree = nodes(writer.root_object.get("/StructTreeRoot"))
    if slug == "titulos":
        for node in tree:
            if node.get("/S") in ("/H1", "/H2"):
                node[NameObject("/S")] = NameObject("/P")
    elif slug == "imagens":
        for node in tree:
            if node.get("/S") == "/Figure":
                node.pop(NameObject("/Alt"), None)
    elif slug == "tabelas":
        for node in tree:
            if node.get("/S") == "/TH":
                node[NameObject("/S")] = NameObject("/TD")
    elif slug == "listas":
        for node in tree:
            if node.get("/S") in ("/L", "/LI", "/Lbl", "/LBody"):
                node[NameObject("/S")] = NameObject("/P")
    elif slug == "metadados":
        writer.add_metadata({"/Title": ""})
        writer.root_object.pop(NameObject("/Lang"), None)
        writer.root_object.pop(NameObject("/Metadata"), None)
    elif slug == "tags":
        for key in ("/StructTreeRoot", "/MarkInfo", "/Metadata"):
            writer.root_object.pop(NameObject(key), None)
    elif slug == "ordem":
        lists = [node for node in tree if node.get("/S") == "/L"]
        assert len(lists) == 1 and len(lists[0]["/K"]) == 2
        lists[0]["/K"].reverse()
    writer.write(target)


def structural_facts(path: Path) -> dict[str, Any]:
    if path.suffix == ".docx":
        doc = Document(path)
        return {"list_items": len(doc.element.xpath(".//w:pPr/w:numPr")),
                "link_text": ["".join(link.itertext()) for link in doc.element.xpath(".//w:hyperlink")]}
    reader = PdfReader(path)
    tree = nodes(reader.root_object.get("/StructTreeRoot"))
    counts = Counter(str(n.get("/S")) for n in tree)
    lists = [n for n in tree if n.get("/S") == "/L"]
    def mcids(value: Any) -> list[int]:
        value = value.get_object() if hasattr(value, "get_object") else value
        if isinstance(value, int):
            return [int(value)]
        if isinstance(value, list):
            return [number for child in value for number in mcids(child)]
        if isinstance(value, dict):
            return mcids(value.get("/K", value.get("/MCID", [])))
        return []
    return {"list_items": counts["/LI"], "tags": dict(counts),
            "list_order_mcids": [[mcids(child) for child in n["/K"]] for n in lists]}


async def build_files(directory: Path) -> dict[str, Any]:
    directory.mkdir(parents=True, exist_ok=True)
    evidence: dict[str, Any] = {"measured_at": datetime.now(timezone.utc).isoformat(), "nvda": "Não executado no Word ou leitor de PDF", "cases": {}}
    images = {"grafico.png": image_bytes()}
    for item in DATA["scenarios"]:
        slug = item["id"]
        spec = spec_for(slug)
        outputs = []
        for fmt in item["formats"]:
            good = directory / f"{slug}-bom.{fmt}"
            bad = directory / f"{slug}-ruim.{fmt}"
            if fmt == "docx":
                documents.build_docx(spec, images, good)
                bad_docx(good, bad, slug)
            else:
                await documents.build_pdf(spec, images, good)
                if slug == "links":
                    changed = spec.model_copy(deep=True)
                    changed.blocks[-1].text = "Clique aqui"
                    await documents.build_pdf(changed, images, bad)
                elif slug == "digitalizacao":
                    page = Image.new("RGB", (1240, 1754), "white")
                    draw = ImageDraw.Draw(page)
                    font = ImageFont.load_default(size=30)
                    for index, text in enumerate(["PDF digitalizado", "Primeiro prepare o documento.", "Depois revise o documento."]):
                        draw.text((100, 150 + 80 * index), text, font=font, fill="black")
                    page.save(bad, "PDF", resolution=150)
                else:
                    bad_pdf(good, bad, slug)
            for kind, path in [("bom", good), ("ruim", bad)]:
                check = await validate(path, external=kind == "bom")
                if kind == "bom":
                    assert check["status"] == "automated_checks_passed", check
                    if fmt == "pdf":
                        assert check["pdf_ua"]["status"] == "passed", check
                elif item["detected"]:
                    assert item["detected"] in {issue["code"] for issue in check["issues"]}, (slug, fmt, check)
                outputs.append({"file": path.name, "kind": kind, "format": fmt, "bytes": path.stat().st_size,
                                "validation": check, "structure": structural_facts(path)})
        evidence["cases"][slug] = outputs
        print(f"{slug}: {len(outputs)} arquivos verificados", flush=True)
    (directory / "medicoes.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return evidence


def render_page(evidence: dict[str, Any], header: str, footer: str) -> str:
    esc = escape
    parts = [header, '<h1>Documentos acessíveis</h1>',
             '<p>Compare arquivos Word DOCX e PDF bons e ruins quanto à propriedade demonstrada. Baixe os arquivos e abra no aplicativo de destino; a página HTML não simula a leitura do documento.</p>',
             '<p>Os defeitos dos arquivos ruins são intencionais. Bom não é certificação de acessibilidade. As medidas abaixo inspecionam os arquivos; a fala real do NVDA no Word ou leitor de PDF ainda não foi capturada.</p>',
             '<nav class="neste-texto" aria-labelledby="neste-texto"><h2 id="neste-texto">Nesta página</h2><ol>' + ''.join(f'<li><a href="#{s["id"]}">{esc(s["title"])}</a></li>' for s in DATA["scenarios"]) + '</ol></nav>']
    for item in DATA["scenarios"]:
        slug = item["id"]
        parts.extend([f'<section aria-labelledby="{slug}"><h2 id="{slug}">{esc(item["title"])}</h2>',
                      f'<h3>O que fazer</h3><p>{esc(item["do"])}</p><h3>O que evitar</h3><p>{esc(item["dont"])}</p>',
                      f'<h3>O que mudou no exemplo ruim</h3><p>{esc(item["defect"])}</p>',
                      '<h3>Arquivos para comparar</h3><ul>'])
        for output in evidence["cases"][slug]:
            label = "Bom" if output["kind"] == "bom" else "Ruim de propósito"
            parts.append(f'<li><a href="documentos/{esc(output["file"])}" download>{label}: {esc(item["title"])} ({output["format"].upper()}, {output["bytes"] // 1024 + 1} KB)</a></li>')
        parts.append('</ul><h3>Estrutura dos arquivos medida pelo MCP</h3>')
        for output in evidence["cases"][slug]:
            check = output["validation"]
            codes = '; '.join(ISSUES.get(issue["code"], "outra propriedade estrutural exige revisão") for issue in check["issues"])
            result = f'Defeitos estruturais encontrados: {codes}.' if codes else 'Nenhum defeito encontrado nas checagens estruturais executadas.'
            external = check.get("pdf_ua", {}).get("status")
            if external == "passed":
                result += ' As verificações automáticas PDF/UA-1 do veraPDF passaram.'
            elif output["format"] == "pdf":
                result += ' veraPDF não executado nesta cópia ruim.'
            if output["kind"] == "ruim" and item["detected"] is None:
                result += ' Este resultado não detecta nem aprova o defeito demonstrado; confira a estrutura específica e o sentido do conteúdo.'
            facts = check["facts"]
            levels = ', '.join(str(level) for level in facts.get("heading_levels", [])) or 'nenhum'
            summary = f'Título: {facts.get("title") or "ausente"}; idioma: {facts.get("language") or "ausente"}; níveis dos títulos: {levels}; itens de lista estruturados: {output["structure"]["list_items"]}.'
            parts.append(f'<p>{esc(output["kind"].capitalize())} {output["format"].upper()}: {esc(result)} {esc(summary)}</p>')
        parts.extend([f'<h3>Comportamento esperado, sem fala capturada</h3><p>{esc(item["expected"])}</p>',
                      f'<h3>Como conferir no aplicativo</h3><p>{esc(item["manual"])}</p>',
                      '<h3>Fontes oficiais</h3><ul>' + ''.join(f'<li><a href="{esc(DATA["sources"][key]["url"])}">{esc(DATA["sources"][key]["title"])}</a></li>' for key in item["sources"]) + '</ul></section>'])
    parts.extend(['<h2>Usar este conhecimento no MCP</h2><p>Peça a criação de um documento ou relatório acessível. O MCP gera os arquivos e devolve validações e limites; o modelo do cliente consulta o guia de comportamento de documentos para escolher a estrutura e explicar as verificações pendentes.</p>',
                  '<h2>Outras verificações necessárias</h2><p>Confira também contraste, informação transmitida só por cor, ampliação e formulários. Esses aspectos não são comprovados por estes pares. OCR exige revisão do texto e da estrutura; PDF com tags ainda exige conferir o sentido e a ordem.</p>',
                  f'<p>Fontes consultadas em {esc(DATA["reviewed"])}. Arquivos medidos em {esc(evidence["measured_at"])}. <a href="documentos/medicoes.json">Baixar o registro técnico das medições</a>.</p>',
                  '<p><a href="index.html">Voltar ao início dos exemplos</a></p>', footer])
    return '\n'.join(parts)


async def main() -> None:
    from construir import CENARIOS, NOME_DO_SITE, RODAPE, cabecalho, pagina_inicial

    evidence = await build_files(HERE / "documentos")
    GUIDE.write_text(guide_text(), encoding="utf-8")
    (HERE / "documentos.html").write_text(render_page(evidence, cabecalho("Documentos acessíveis | Exemplos", "documentos"), RODAPE), encoding="utf-8")
    # So a moldura muda nas paginas web; suas medicoes permanecem intactas.
    start = '<main id="conteudo" tabindex="-1">\n<div class="coluna">\n'
    end = '</div>\n</main>\n<footer>'
    for scenario in CENARIOS:
        path = HERE / f'{scenario["slug"]}.html'
        if path.is_file():
            original = path.read_text(encoding="utf-8")
            assert start in original and end in original, path
            body = original.partition(start)[2].rpartition(end)[0]
            path.write_text(cabecalho(f'{scenario["nome"]} | {NOME_DO_SITE}', scenario["slug"]) + body + RODAPE, encoding="utf-8")
    (HERE / "index.html").write_text(pagina_inicial(), encoding="utf-8")


if __name__ == "__main__":
    asyncio.run(main())
