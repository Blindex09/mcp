"""Arquivos reais, defeitos plantados e fronteiras de seguranca dos documentos."""

import json
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image
from pydantic import ValidationError
from pypdf import PdfReader, PdfWriter

from accessibility_mcp.a11y import documents, files
from accessibility_mcp.a11y.document_schema import Assessment, Block, DocumentSpec
from accessibility_mcp.a11y.document_validation import (
    validate,
    validate_docx,
    validate_pdf,
)


@pytest.fixture(autouse=True)
def no_validator_download_in_unit_tests(monkeypatch):
    monkeypatch.setenv("A11Y_MCP_AUTO_INSTALL", "0")


@pytest.fixture
def spec():
    return DocumentSpec(title="Relatório de acessibilidade", author="Equipe de QA", blocks=[
        Block(kind="paragraph", text="Avaliação de teclado e formulários com evidências."),
        Block(kind="heading", text="Resultados"),
        Block(kind="table", text="Critérios avaliados", headers=["Critério", "Resultado"], rows=[["2.1.1", "Não avaliado"]]),
        Block(kind="list", items=["Conferir o foco", "Verificar os anúncios"]),
        Block(kind="link", text="Diretrizes do W3C", url="https://www.w3.org/WAI/"),
    ])


@pytest.mark.parametrize("status", ["Supports", "Not Applicable", "Partially Supports", "Does Not Support"])
def test_conformance_needs_evidence_not_just_a_status(status):
    with pytest.raises(ValidationError, match="evidencia"):
        Assessment(criterion="2.1.1", status=status, coverage="complete", remarks="Sem achados")


def test_positive_claim_needs_complete_evaluation_and_valid_criterion():
    with pytest.raises(ValidationError, match="cobertura"):
        Assessment(criterion="2.1.1", status="Supports", evidence=["Scanner verde"], remarks="Teste parcial")
    with pytest.raises(ValidationError, match="inexistente"):
        Assessment(criterion="9.9.9", status="Not Evaluated", remarks="Sem teste")
    assert Assessment(criterion="2.1.1", status="Not Evaluated", remarks="Falta teclado").coverage == "not_evaluated"


@pytest.mark.parametrize("block", [
    {"kind": "heading", "level": 3, "text": "Pulo"},
    {"kind": "table", "headers": ["Nome"], "rows": [["A", "B"]]},
    {"kind": "image", "source": "foto.png"},
    {"kind": "link", "text": "Executar", "url": "javascript:alert(1)"},
    {"kind": "link", "text": "Site", "url": "http:sem-servidor"},
    {"kind": "paragraph", "text": " "},
])
def test_incomplete_or_unsafe_document_is_rejected_before_writing(block):
    with pytest.raises(ValidationError):
        DocumentSpec.model_validate({"title": "Teste", "blocks": [block]})


async def test_docx_is_an_actual_semantic_document_and_mutations_are_caught(spec):
    generated = await documents.create(spec, "docx")
    path = Path(generated["files"][0]["path"])
    result = validate_docx(path)
    assert result["issues"] == []
    assert result["facts"]["heading_levels"] == [1, 2]
    assert result["facts"]["language"] == "pt-BR"
    assert result["conformance"] == "not_determined" and not result["complete_validation"]
    doc = Document(str(path))
    assert doc.core_properties.title == spec.title
    assert not doc.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    assert doc.tables[0].cell(1, 1).text == "Não avaliado"
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in doc.tables[0].rows)
    doc.core_properties.title = ""
    row_properties = doc.tables[0].rows[0]._tr.trPr
    row_properties.remove(row_properties[0])
    doc.save(str(path))
    assert {i["code"] for i in validate_docx(path)["issues"]} >= {"missing_title", "table_missing_header"}


async def test_pdf_has_tags_unicode_fonts_headings_tables_lists_and_links(spec):
    generated = await documents.create(spec, "pdf")
    result = generated["files"][0]["validation"]
    assert result["issues"] == []
    assert result["facts"]["tags"]["/TH"] == 2
    assert result["facts"]["tags"]["/L"] == 1 and result["facts"]["tags"]["/Link"] == 1
    path = Path(generated["files"][0]["path"])
    reader = PdfReader(path)
    assert "Não avaliado" in reader.pages[0].extract_text()
    assert b"<pdfuaid:part>1</pdfuaid:part>" in reader.root_object["/Metadata"].get_data()
    assert reader.pages[0]["/Tabs"] == "/S"
    assert reader.pages[0]["/Annots"][0].get_object()["/Contents"] == "Diretrizes do W3C"
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    del writer.root_object["/StructTreeRoot"]
    writer.write(path)
    assert "missing_tags" in {i["code"] for i in validate_pdf(path)["issues"]}


async def test_pdf_link_description_uses_the_url_actually_normalized_by_the_browser(spec):
    spec.blocks = [Block(kind="link", text="Referência do W3C", url="https://WWW.W3.ORG")]
    result = await documents.create(spec, "pdf")
    reader = PdfReader(result["files"][0]["path"])
    annotation = reader.pages[0]["/Annots"][0].get_object()
    assert annotation["/A"]["/URI"] == "https://www.w3.org/"
    assert annotation["/Contents"] == "Referência do W3C"


async def test_local_image_alt_survives_both_formats_and_missing_alt_is_detected(spec, tmp_path, monkeypatch):
    monkeypatch.setattr(files, "_CHOSEN", tmp_path)
    Image.new("RGB", (300, 100), "white").save(tmp_path / "imagem.png")
    spec.blocks.append(Block(kind="image", source="imagem.png", alt="Fluxo de avaliação do formulário"))
    result = await documents.create(spec)
    for entry in result["files"]:
        assert not entry["validation"]["issues"]
    docx_path = Path(result["files"][0]["path"])
    doc = Document(str(docx_path))
    assert doc.inline_shapes[0]._inline.docPr.get("descr") == "Fluxo de avaliação do formulário"
    del doc.inline_shapes[0]._inline.docPr.attrib["descr"]
    doc.save(str(docx_path))
    assert "image_missing_alt" in {i["code"] for i in validate_docx(docx_path)["issues"]}


async def test_tall_image_fits_the_docx_page(spec, tmp_path, monkeypatch):
    from docx.shared import Cm

    monkeypatch.setattr(files, "_CHOSEN", tmp_path)
    Image.new("RGB", (50, 1000), "white").save(tmp_path / "alta.png")
    spec.blocks.append(Block(kind="image", source="alta.png", alt="Figura vertical"))
    result = await documents.create(spec, "docx")
    shape = Document(result["files"][0]["path"]).inline_shapes[0]
    assert shape.width <= Cm(12) and shape.height <= Cm(20)


def test_text_is_escaped_and_no_untrusted_html_can_run(spec):
    spec.blocks = [Block(kind="paragraph", text='<script src="http://localhost/segredo"></script>')]
    html = documents.build_html(spec, {})
    assert "<script" not in html and "&lt;script" in html
    assert '<html lang="pt-BR">' in html


def test_document_paths_cannot_escape_or_read_secret_files(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "_CHOSEN", tmp_path)
    for document_id, filename in [("../fora", "document.pdf"), ("a" * 32, "../../.env"), ("a" * 32, "validation.json")]:
        with pytest.raises(ValueError):
            documents.artifact_path(document_id, filename)
    with pytest.raises(files.FileError):
        documents.resolve_input(str(tmp_path / ".env"))
    with pytest.raises(files.FileError):
        documents.resolve_input(str(tmp_path.parent / "fora.pdf"))


async def test_outputs_do_not_overwrite_and_can_be_read_as_resources(spec):
    from accessibility_mcp.server import mcp

    first = await documents.create(spec, "docx")
    second = await documents.create(spec, "docx")
    assert first["document_id"] != second["document_id"]
    contents = await mcp.read_resource(first["files"][0]["uri"])
    assert next(iter(contents)).content[:2] == b"PK"


def test_empty_report_never_invents_conformance_or_passes():
    result = documents.report_spec({"url": "https://exemplo.test", "findings": []}, "Avaliação")
    assert result.assessments == []
    text = " ".join(b.text for b in result.blocks)
    assert "não comprova conformidade" in text and "não avaliados" in text
    with pytest.raises(ValueError, match="portugues"):
        documents.report_spec({}, "Report", "en")


async def test_report_from_saved_run_keeps_evidence_and_coverage():
    from accessibility_mcp.a11y import runs
    from accessibility_mcp.server import mcp

    run_id = runs.save("exemplo.test", {"url": "https://exemplo.test", "findings": [{"criterion": "2.1.1", "claim": "Controle sem teclado", "outcome": "failed", "evidence_steps": [3]}], "coverage": {"gaps": ["NVDA não executado"]}})
    tool = mcp._tool_manager.get_tool("a11y_create_report")
    result = json.loads(await tool.fn(formats="docx", run_id=run_id))
    assert "error" not in result
    path = Path(result["files"][0]["path"])
    text = " ".join(p.text for p in Document(str(path)).paragraphs)
    assert "Controle sem teclado" in text and "[3]" in text and "NVDA não executado" in text


async def test_external_validator_missing_is_not_a_pass(spec, monkeypatch):
    from accessibility_mcp.a11y import document_validation

    monkeypatch.delenv("A11Y_MCP_VERAPDF", raising=False)
    monkeypatch.setattr(document_validation.shutil, "which", lambda _: None)
    result = await documents.create(spec, "pdf")
    check = await validate(Path(result["files"][0]["path"]))
    assert check["pdf_ua"]["status"] == "not_run"
    assert check["status"] == "partial"
    assert check["conformance"] == "not_determined"


async def test_pdf_empty_title_is_missing_instead_of_the_string_none(spec, tmp_path):
    from pypdf import PdfWriter

    result = await documents.create(spec, "pdf")
    writer = PdfWriter(clone_from=result["files"][0]["path"])
    writer.add_metadata({"/Title": ""})
    path = tmp_path / "sem-titulo.pdf"
    writer.write(path)
    check = validate_pdf(path)
    assert check["facts"]["title"] == ""
    assert "missing_title" in {issue["code"] for issue in check["issues"]}


def test_docx_with_external_entities_is_rejected(spec, tmp_path):
    from defusedxml.common import DefusedXmlException

    path = tmp_path / "hostil.docx"
    documents.build_docx(spec, {}, path)
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    members["word/document.xml"] = b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///segredo">]><x>&xxe;</x>'
    with ZipFile(path, "w") as target:
        for name, data in members.items():
            target.writestr(name, data)
    with pytest.raises(DefusedXmlException):
        validate_docx(path)


async def test_verapdf_installer_hash_mismatch_never_executes_code(monkeypatch):
    import httpx

    from accessibility_mcp.a11y import verapdf_runtime

    monkeypatch.setenv("A11Y_MCP_AUTO_INSTALL", "1")
    monkeypatch.setattr(verapdf_runtime.shutil, "which", lambda _: sys.executable)
    original_client = httpx.AsyncClient
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=b"instalador adulterado"))
    monkeypatch.setattr(verapdf_runtime.httpx, "AsyncClient", lambda **kwargs: original_client(transport=transport, **kwargs))
    with pytest.raises(ValueError, match="hash"):
        await verapdf_runtime.ensure()
    assert not verapdf_runtime.cached_jar().exists()


@pytest.mark.parametrize(("xml", "returncode", "expected"), [
    (b'<report><validationReport isCompliant="true"/></report>', 0, "passed"),
    (b'<report><validationReport isCompliant="true"/></report>', 2, "error"),
    (b'<report><validationReport isCompliant="false"/></report>', 1, "failed"),
    (b'<report><batchSummary/></report>', 0, "error"),
    (b'nao e XML', 0, "error"),
])
async def test_verapdf_exit_and_report_are_both_checked(monkeypatch, tmp_path, xml, returncode, expected):
    from accessibility_mcp.a11y import document_validation

    class Process:
        def __init__(self):
            self.returncode = returncode

        async def communicate(self):
            return xml, b""

    async def launch(*args, **kwargs):
        assert "--flavour" in args and "ua1" in args
        return Process()

    monkeypatch.setenv("A11Y_MCP_VERAPDF", sys.executable)
    monkeypatch.setattr(document_validation.asyncio, "create_subprocess_exec", launch)
    assert (await document_validation.verapdf(tmp_path / "arquivo.pdf"))["status"] == expected
